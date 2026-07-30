import errno
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation

from bench import f0_storage


LOGICAL = "a" * 64
PHYSICAL = "00000000-0000-4000-8000-000000000001"


@pytest.fixture
def office_payloads(tmp_path):
    return f0_storage.create_office_templates(tmp_path / "templates")


def test_marker_last_candidate_is_hash_valid_and_office_reopenable(
    tmp_path, office_payloads
):
    candidate = f0_storage.prepare_candidate(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
        office_payloads,
    )
    assert f0_storage.classify_candidate(candidate) == "committed"
    assert (candidate / f0_storage.COMMITTED).read_bytes() == b""
    assert (
        f0_storage.validate_committed(candidate)["logical_hash"] == LOGICAL
    )

    load_workbook(
        str(candidate / "artifacts" / "probe.xlsx"), read_only=True
    ).close()
    assert len(
        Presentation(str(candidate / "artifacts" / "probe.pptx")).slides
    ) == 1


def test_prepared_candidate_is_adopted_without_rewriting_evidence(
    tmp_path, office_payloads
):
    candidate = f0_storage.prepare_candidate(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
        office_payloads,
        commit=False,
    )
    before = {
        name: (candidate / name).read_bytes()
        for name in f0_storage.REQUIRED_FILES
    }
    assert f0_storage.classify_candidate(candidate) == "prepared"
    assert f0_storage.recover_candidate(candidate) == "committed"
    assert before == {
        name: (candidate / name).read_bytes()
        for name in f0_storage.REQUIRED_FILES
    }


@pytest.mark.parametrize("boundary", f0_storage.BOUNDARIES)
def test_every_in_process_fault_boundary_fails_closed(
    tmp_path, office_payloads, boundary
):
    root = tmp_path / boundary
    candidate = f0_storage.candidate_path(
        root, LOGICAL, PHYSICAL
    )
    with pytest.raises(f0_storage.InjectedStop, match=boundary):
        f0_storage.prepare_candidate(
            root,
            LOGICAL,
            PHYSICAL,
            office_payloads,
            crash_after=boundary,
        )
    state = f0_storage.classify_candidate(candidate)
    if state == "prepared":
        assert f0_storage.recover_candidate(candidate) == "committed"
    elif state == "committed":
        f0_storage.validate_committed(candidate)
    else:
        assert state == "abandoned"
        assert not (candidate / f0_storage.COMMITTED).exists()


def test_hash_tamper_and_unexpected_members_invalidate_commit(
    tmp_path, office_payloads
):
    first = f0_storage.prepare_candidate(
        tmp_path / "first", LOGICAL, PHYSICAL, office_payloads
    )
    with (first / "artifacts" / "probe.xlsx").open("ab") as handle:
        handle.write(b"tamper")
    assert f0_storage.classify_candidate(first) == "corrupt_committed"
    with pytest.raises(f0_storage.StorageIntegrityError, match="mismatch"):
        f0_storage.validate_committed(first)

    second = f0_storage.prepare_candidate(
        tmp_path / "second", LOGICAL, PHYSICAL, office_payloads
    )
    (second / "unexpected.txt").write_text("unexpected", encoding="utf-8")
    assert f0_storage.classify_candidate(second) == "corrupt_committed"


def test_physical_directory_collision_never_overwrites(
    tmp_path, office_payloads
):
    root = tmp_path / "store"
    candidate = f0_storage.prepare_candidate(
        root, LOGICAL, PHYSICAL, office_payloads
    )
    before = (candidate / "attempt.json").read_bytes()
    with pytest.raises(FileExistsError):
        f0_storage.prepare_candidate(
            root, LOGICAL, PHYSICAL, office_payloads
        )
    assert (candidate / "attempt.json").read_bytes() == before
    f0_storage.validate_committed(candidate)


def test_retry_policy_retries_only_retryable_filesystem_errors(
    monkeypatch, tmp_path, office_payloads
):
    candidate = f0_storage.prepare_candidate(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
        office_payloads,
        commit=False,
    )
    real_validate = f0_storage.validate_prepared
    calls = {"count": 0}

    def flaky(path):
        calls["count"] += 1
        if calls["count"] < 3:
            raise OSError(errno.EACCES, "sharing simulation")
        return real_validate(path)

    elapsed = {"value": 0.0}
    monkeypatch.setattr(f0_storage, "validate_prepared", flaky)
    assert f0_storage.publish_prepared(
        candidate,
        clock=lambda: elapsed["value"],
        sleeper=lambda delay: elapsed.__setitem__(
            "value", elapsed["value"] + delay
        ),
    ) == "committed"
    assert calls["count"] >= 3


def test_publication_deadline_never_creates_a_marker(
    monkeypatch, tmp_path, office_payloads
):
    candidate = f0_storage.prepare_candidate(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
        office_payloads,
        commit=False,
    )
    monkeypatch.setattr(
        f0_storage,
        "validate_prepared",
        lambda _path: (_ for _ in ()).throw(
            OSError(errno.EACCES, "persistent sharing violation")
        ),
    )
    elapsed = {"value": 0.0}
    assert (
        f0_storage.publish_prepared(
            candidate,
            deadline_seconds=0.1,
            clock=lambda: elapsed["value"],
            sleeper=lambda delay: elapsed.__setitem__(
                "value", elapsed["value"] + delay
            ),
        )
        == "publish_blocked"
    )
    assert not (candidate / f0_storage.COMMITTED).exists()


def test_nonretryable_publication_error_is_not_retried(
    monkeypatch, tmp_path, office_payloads
):
    candidate = f0_storage.prepare_candidate(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
        office_payloads,
        commit=False,
    )
    calls = {"count": 0}

    def invalid(_path):
        calls["count"] += 1
        raise OSError(errno.EINVAL, "not retryable")

    monkeypatch.setattr(f0_storage, "validate_prepared", invalid)
    with pytest.raises(OSError):
        f0_storage.publish_prepared(candidate)
    assert calls["count"] == 1
    assert not (candidate / f0_storage.COMMITTED).exists()


def test_normal_candidate_routes_marker_creation_through_retry(
    monkeypatch, tmp_path, office_payloads
):
    real_create = f0_storage.create_commit_marker
    calls = {"count": 0}

    def transient(candidate):
        calls["count"] += 1
        if calls["count"] == 1:
            raise OSError(errno.EACCES, "Defender simulation")
        return real_create(candidate)

    monkeypatch.setattr(f0_storage, "create_commit_marker", transient)
    candidate = f0_storage.prepare_candidate(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
        office_payloads,
    )
    assert calls["count"] == 2
    f0_storage.validate_committed(candidate)


def test_normal_candidate_retries_first_prepared_reread(
    monkeypatch, tmp_path, office_payloads
):
    real_validate = f0_storage.validate_prepared
    calls = {"count": 0}

    def transient(candidate):
        calls["count"] += 1
        if calls["count"] == 1:
            raise OSError(errno.EACCES, "indexer simulation")
        return real_validate(candidate)

    elapsed = {"value": 0.0}
    monkeypatch.setattr(f0_storage, "validate_prepared", transient)
    candidate = f0_storage.prepare_candidate(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
        office_payloads,
        clock=lambda: elapsed["value"],
        sleeper=lambda delay: elapsed.__setitem__(
            "value", elapsed["value"] + delay
        ),
    )
    assert calls["count"] >= 2
    f0_storage.validate_committed(candidate)


def test_first_prepared_reread_obeys_publication_deadline(
    monkeypatch, tmp_path, office_payloads
):
    monkeypatch.setattr(
        f0_storage,
        "validate_prepared",
        lambda _path: (_ for _ in ()).throw(
            OSError(errno.EACCES, "persistent indexer simulation")
        ),
    )
    elapsed = {"value": 0.0}
    with pytest.raises(
        f0_storage.StorageIntegrityError,
        match="publish_blocked",
    ):
        f0_storage.prepare_candidate(
            tmp_path / "store",
            LOGICAL,
            PHYSICAL,
            office_payloads,
            deadline_seconds=0.1,
            clock=lambda: elapsed["value"],
            sleeper=lambda delay: elapsed.__setitem__(
                "value", elapsed["value"] + delay
            ),
        )
    candidate = f0_storage.candidate_path(
        tmp_path / "store",
        LOGICAL,
        PHYSICAL,
    )
    assert elapsed["value"] == pytest.approx(0.1)
    assert not (candidate / f0_storage.COMMITTED).exists()


def test_small_hard_exit_spike_has_no_invalid_visible_bundle(tmp_path):
    summary = f0_storage.run_spike(
        tmp_path / "spike",
        cycles=len(f0_storage.BOUNDARIES) + 1,
        crash_cycles=len(f0_storage.BOUNDARIES),
    )
    assert summary["passed"]
    assert summary["cycles"] == len(f0_storage.BOUNDARIES) + 1
    assert summary["forced_exits"] == len(f0_storage.BOUNDARIES)
    assert summary["invalid_committed"] == 0
    assert summary["directory_renames"] == 0
    assert summary["committed"] == summary["cycles"]
    assert summary["logical_commits"] == summary["cycles"]
    assert summary["duplicate_valid_candidates"] == {}
    assert summary["physical_candidates"] >= summary["cycles"]


def test_storage_probe_source_never_renames_or_replaces_directories():
    source = Path(f0_storage.__file__).read_text(encoding="utf-8")
    assert "os.rename" not in source
    assert "os.replace" not in source
    assert "Path.rename" not in source
    assert "Path.replace" not in source


@pytest.mark.skipif(os.name != "nt", reason="Windows held-handle smoke test")
def test_windows_real_office_handle_uses_bounded_publication(tmp_path):
    summary = f0_storage.run_spike(
        tmp_path / "windows-held-handle",
        cycles=1,
        crash_cycles=0,
        held_handle_cycles=1,
    )
    assert summary["passed"]
    assert summary["held_handle_cycles"] == 1
    assert summary["records"][0]["publish_elapsed_seconds"] >= 0.1
